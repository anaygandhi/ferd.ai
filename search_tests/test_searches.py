from fpdf import FPDF
import os
import sys
from docx import Document
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from pdf_search import search_in_pdf
from pdf_search import search_in_docx

# Creating a test PDF file
def create_test_pdf(filename="test.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Hello, this is a test PDF.", ln=True, align='C')
    pdf.output(filename)

# Testing the PDF search functionality
def test_pdf_functionality():
    test_pdf = "test.pdf"
    create_test_pdf(test_pdf)
    expected_text = "Hello, this is a test PDF."
    
    if os.path.exists(test_pdf):
        extracted_text = search_in_pdf(test_pdf).strip()
        print("\nExtracted Text from PDF:\n", extracted_text)

        if expected_text not in extracted_text:
            print(f"Test Failed: Expected '{expected_text}', but got '{extracted_text}'")
        else:
            print("Test Passed: Extracted text matches expected content")
    else:
        print(f"Test Failed: {test_pdf} not found")

# Creating a test docx file
def create_test_docx(filename="test.docx"):
    doc = Document()
    doc.add_paragraph("Hello, this is a test docx file.")
    doc.save(filename)

# Testing the docx search functionality
def test_docx_functionality():
    test_docx = "test.docx"
    create_test_docx(test_docx)
    expected_text = "Hello, this is a test docx file."
    
    if os.path.exists(test_docx):
        extracted_text = search_in_docx(test_docx).strip()
        print("\nExtracted Text from DOCX:\n", extracted_text)

        if expected_text not in extracted_text:
            print(f"Test Failed: Expected '{expected_text}', but got '{extracted_text}'")
        else:
            print("Test Passed: Extracted text matches expected content")
    else:
        print(f"Test Failed: {test_docx} not found")

# Run  
if __name__ == "__main__":
    test_pdf_functionality()
    test_docx_functionality()
